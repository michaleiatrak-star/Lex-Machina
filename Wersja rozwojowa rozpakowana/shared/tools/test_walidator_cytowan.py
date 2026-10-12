#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Testy walidator_cytowan.py: domena urzędowa po hoście, pełny identyfikator cytatu,
kontrola zapisu powołań (shared/FORMAT-CYTOWANIA.md, zgł. #109).

Uruchomienie:  python3 -m unittest test_walidator_cytowan -v   (z katalogu shared/tools)
"""
import importlib.util
import os
import subprocess
import sys
import tempfile
import unittest

_spec = importlib.util.spec_from_file_location(
    "walidator_cytowan", os.path.join(os.path.dirname(os.path.abspath(__file__)), "walidator_cytowan.py"))
w = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(w)


def potwierdzone(cytat, **zdarzenie):
    c = w.extract_citations(cytat)[0]  # typ cytatu (artykul / dziennik_ustaw / sygnatura) z ekstraktora
    return w.log_has_verification(c, [zdarzenie])[0]


class DomenaUrzedowa(unittest.TestCase):
    def test_podciag_domeny_nie_wystarcza(self):
        for url in ("https://msn.pl/211", "https://sn.pl.evil.example/211",
                    "https://evil.example/isap.sejm.gov.pl/211", "https://evil.example/?u=sn.pl&a=211"):
            self.assertFalse(potwierdzone("art. 211 KC", url=url, query_context="art. 211"), url)

    def test_wymagane_https(self):
        self.assertFalse(potwierdzone("art. 211 KC", url="http://isap.sejm.gov.pl/x", query_context="art. 211"))

    def test_subdomena_i_wynik_wyszukiwania(self):
        self.assertTrue(potwierdzone("I CSK 123/24", tool="web_search", query="I CSK 123/24",
                                     result_urls=["https://www.sn.pl/sites/orzecznictwo/I CSK 123-24.pdf"]))
        self.assertFalse(potwierdzone("I CSK 123/24", tool="web_search", query="I CSK 123/24",
                                      result_urls=["https://www.sn.pl.attacker.net/x"]))


class ZgodnoscCytatu(unittest.TestCase):
    def test_cala_liczba_nie_podciag(self):
        self.assertFalse(potwierdzone("art. 211 KC", url="https://isap.sejm.gov.pl/x", query_context="art. 2110 kc"))
        self.assertTrue(potwierdzone("art. 211 KC", url="https://isap.sejm.gov.pl/x", query_context="art. 211 kc"))

    def test_wszystkie_liczby_cytatu(self):
        self.assertFalse(potwierdzone("Dz.U. 2023 poz. 1691", url="https://isap.sejm.gov.pl/a", query_context="ustawa 2023"))
        self.assertFalse(potwierdzone("I CSK 123/24", url="https://orzeczenia.nsa.gov.pl/doc/X", query_context="I CSK 1234/24"))

    def test_sygnatura_rozna_litera_izby(self):
        self.assertFalse(potwierdzone("I CSK 123/24", url="https://www.sn.pl/x", query_context="II CSK 123/24"))

    def test_identyfikator_isap(self):
        self.assertTrue(potwierdzone("Dz.U. 2023 poz. 1691",
                                     url="https://isap.sejm.gov.pl/isap.nsf/DocDetails.xsp?id=WDU20230001691"))


def kody(tekst):
    return {n["kod"] for n in w.kontrola_zapisu(tekst)}


class EkstrakcjaPowolan(unittest.TestCase):
    def test_artykul_ze_skrotem_z_kropkami_i_jednostkami(self):
        for cytat in ("art. 130 § 1 k.p.c.", "art. 398²² k.p.c.", "art. 9 ust. 1 u.KRS", "art. 415 § 1 pkt 2 k.c."):
            self.assertEqual([c["tekst"] for c in w.extract_citations(cytat)], [cytat])

    def test_warianty_dziennika(self):
        for cytat in ("Dz.U.2025.383", "Dz. U. poz. 1061", "Dz. U. z 1997 r. Nr 89, poz. 555"):
            self.assertEqual([c["tekst"] for c in w.extract_citations(cytat)], [cytat])

    def test_skrot_kodeksu_w_logu(self):
        self.assertTrue(potwierdzone("art. 130 § 1 k.p.c.", url="https://isap.sejm.gov.pl/x", query_context="art. 130 § 1 KPC"))
        self.assertFalse(potwierdzone("art. 130 k.p.c.", url="https://isap.sejm.gov.pl/x", query_context="art. 130 k.p.k."))

    def test_indeks_gorny_w_logu(self):
        self.assertTrue(potwierdzone("art. 398²² k.p.c.", url="https://isap.sejm.gov.pl/x", query_context="art. 398^22 KPC"))
        self.assertFalse(potwierdzone("art. 398²² k.p.c.", url="https://isap.sejm.gov.pl/x", query_context="art. 398 KPC"))


class KontrolaZapisu(unittest.TestCase):
    def test_bledne_oznaczenia_dziennika(self):
        self.assertEqual(kody("(t.j. Dz.U. z 2025 r. poz. 869 ze zm.)"), {"F1", "F3", "F4"})
        self.assertEqual(kody("(Dz.U. 2025 poz. 869)"), {"F1", "F2"})
        self.assertEqual(kody("(Dz.U.2025.383 t.j.)"), {"F1", "F3", "F5"})
        self.assertEqual(kody("(Dz. U. z 2025 poz. 869)"), {"F2"})

    def test_poprawne_oznaczenia_dziennika(self):
        for tekst in ("(Dz. U. z 2025 r. poz. 869)", "(Dz. U. z 2025 r. poz. 869 z późn. zm.)",
                      "(Dz. U. z 1997 r. Nr 89, poz. 555)", "(Dz. U. poz. 1061)",
                      "(Dz. U. z 2025 r. poz. 869, z 2026 r. poz. 12 i 40)"):
            self.assertEqual(kody(tekst), set(), tekst)

    def test_jednostki_i_skroty_kodeksow(self):
        self.assertEqual(kody("art. 130 §1 KPC"), {"F7", "F8"})
        self.assertEqual(kody("art. 5, ust. 2 ustawy"), {"F8"})
        self.assertEqual(kody("kodeks (Dz. U. z 2025 r. poz. 1; dalej: k.p.c.), art. 398^22 § 1 k.p.c."), {"F8"})

    def test_skrot_przed_wprowadzeniem_i_niewprowadzony(self):
        self.assertEqual(kody("art. 5 k.c. … (Dz. U. z 2026 r. poz. 795; dalej: k.c.)"), {"F9"})
        self.assertEqual(kody("art. 5 k.c."), {"F9"})
        self.assertEqual(kody("(Dz. U. z 2026 r. poz. 795; dalej: k.c.), potem art. 415 k.c."), set())
        self.assertEqual(kody("(Dz. U. z 2026 r. poz. 795), zwana dalej „k.c.”, potem art. 415 k.c."), set())

    def test_metryka_tylko_w_czesci_roboczej(self):
        self.assertEqual(kody("(DU/2025/869, IN_FORCE)"), {"F6"})
        self.assertEqual(kody("Źródło: https://isap.sejm.gov.pl/isap.nsf/DocDetails.xsp?id=WDU20250000869"), {"F6"})
        self.assertEqual(kody("Treść.\n\n## CZĘŚĆ ROBOCZA\n(DU/2025/869, IN_FORCE) Dz.U. 2025 poz. 869"), set())

    def test_tj_poza_powolaniem_nie_jest_naruszeniem(self):
        self.assertEqual(kody("Sąd stosuje t.j. ogłoszony w marcu."), set())

    def test_cli_kod_wyjscia_3(self):
        skrypt = os.path.join(os.path.dirname(os.path.abspath(__file__)), "walidator_cytowan.py")
        with tempfile.TemporaryDirectory() as tmp:
            zle, dobrze = os.path.join(tmp, "zle.md"), os.path.join(tmp, "dobrze.md")
            with open(zle, "w", encoding="utf-8") as f:
                f.write("Na podstawie art. 130 §1 KPC (Dz.U. 2025 poz. 869).")
            with open(dobrze, "w", encoding="utf-8") as f:
                f.write("Ustawa (Dz. U. z 2025 r. poz. 869; dalej: k.p.c.); art. 130 § 1 k.p.c.")
            run = lambda doc: subprocess.run([sys.executable, skrypt, "--document", doc, "--format"],
                                            capture_output=True, text=True)
            self.assertEqual(run(zle).returncode, 3)
            self.assertEqual(run(dobrze).returncode, 0)

    def test_cli_kod_wyjscia_3_konsola_cp1250(self):
        # Konsola Windows: wyjście w cp1250 bez „✗”; kod nadal 3, nie 1 (UnicodeEncodeError, #109).
        skrypt = os.path.join(os.path.dirname(os.path.abspath(__file__)), "walidator_cytowan.py")
        with tempfile.TemporaryDirectory() as tmp:
            zle = os.path.join(tmp, "zle.md")
            with open(zle, "w", encoding="utf-8") as f:
                f.write("Na podstawie art. 130 §1 KPC (Dz.U. 2025 poz. 869).")
            env = {**os.environ, "PYTHONIOENCODING": "cp1250"}
            run = subprocess.run([sys.executable, skrypt, "--document", zle, "--format"],
                                 capture_output=True, env=env)
            self.assertEqual(run.returncode, 3, run.stderr.decode("utf-8", "replace"))


try:
    import docx  # python-docx
except ImportError:
    docx = None


@unittest.skipIf(docx is None, "python-docx niedostępny")
class OdczytDocx(unittest.TestCase):
    def test_tabela_i_indeks_gorny(self):
        d = docx.Document()
        p = d.add_paragraph("Na podstawie art. 398")
        sup = p.add_run("22")
        sup.font.superscript = True
        p.add_run(" k.p.c. (Dz. U. z 2025 r. poz. 1; dalej: k.p.c.)")
        d.add_table(rows=1, cols=1).cell(0, 0).text = "art. 130 §1 KPC"
        with tempfile.TemporaryDirectory() as tmp:
            sciezka = os.path.join(tmp, "pismo.docx")
            d.save(sciezka)
            from pathlib import Path
            tekst = w.read_document_text(Path(sciezka))
        self.assertIn("art. 398²² k.p.c.", tekst)
        self.assertIn("art. 130 §1 KPC", tekst)
        self.assertIn("F7", kody(tekst))


if __name__ == "__main__":
    unittest.main()
